from docx import Document 
from docx.shared import Inches 
 
def create_report(): 
    doc = Document() 
 
    # Title 
    doc.add_heading('AI-Based Crop Recommendation and Health Monitoring System', 0) 
 
    # Abstract 
    doc.add_heading('Abstract', level=1) 
    doc.add_paragraph( 
        'This project presents an intelligent farming assistant named "Kisan Sahayak". ' 
        'The system leverages Machine Learning (ML) and Artificial Intelligence (AI) to provide ' 
        'real-time crop recommendations, predict potential farming diseases based on weather patterns, ' 
        'and analyze plant health using leaf images. Additionally, the system provides both traditional ' 
        'Ayurvedic remedies and modern commercial medicine suggestions for identified crop issues.' 
    ) 
 
    # Objectives 
    doc.add_heading('Objectives', level=1) 
    objectives = [ 
        'To provide data-driven crop recommendations based on soil parameters (N, P, K, pH) and weather.', 
        'To predict potential diseases and pests using current weather conditions.', 
        'To enable visual plant health monitoring through deep learning-based image analysis.', 
        'To offer holistic treatment options, including natural/Ayurvedic and commercial solutions.' 
    ] 
    for obj in objectives: 
        doc.add_paragraph(obj, style='List Bullet') 
 
    # Methodology 
    doc.add_heading('Methodology', level=1) 
    doc.add_paragraph( 
        'The system is built using the Flask web framework. Three distinct ML/DL models are integrated:' 
    ) 
    doc.add_paragraph('1. Crop Recommendation: Random Forest Classifier trained on soil and weather data.', style='List Number') 
    doc.add_paragraph('2. Disease Prediction: Random Forest Classifier trained on historical weather-disease patterns.', style='List Number') 
    doc.add_paragraph('3. Image Analysis: Convolutional Neural Network (CNN) for leaf disease classification.', style='List Number') 
 
    # Features 
    doc.add_heading('Key Features', level=1) 
    features = [ 
        'Real-time Soil Analysis Integration', 
        'Weather-based Early Warning System', 
        'Visual Health Monitoring (Image Upload)', 
        'Comprehensive Treatment Database (Ayurvedic + Commercial + Prevention)', 
        'Multilingual UI Support (Hindi & English)',
        'Architecture Diagram and Modular Design',
        'Modern, Responsive UI built with Bootstrap' 
    ] 
    for feat in features: 
        doc.add_paragraph(feat, style='List Bullet') 

    # Requirements
    doc.add_heading('System Requirements', level=1)
    doc.add_heading('Hardware', level=2)
    hw = ['Processor: Intel i3 or better', 'RAM: 4GB minimum', 'Storage: 500MB', 'Camera for Image Capture']
    for item in hw:
        doc.add_paragraph(item, style='List Bullet')
    
    doc.add_heading('Software', level=2)
    sw = ['Python 3.8+', 'Flask', 'TensorFlow & Scikit-learn', 'Pandas & NumPy']
    for item in sw:
        doc.add_paragraph(item, style='List Bullet')

    # Future Scope
    doc.add_heading('Future Scope', level=1)
    scope = [
        'IoT Sensor Integration for real-time soil data.',
        'Mobile Application for Android and iOS.',
        'Market integration for direct selling.',
        'Voice Assistant for illiterate farmers.'
    ]
    for item in scope:
        doc.add_paragraph(item, style='List Bullet')
 
    # Conclusion 
    doc.add_heading('Conclusion', level=1) 
    doc.add_paragraph( 
        'Kisan Sahayak bridge the gap between traditional agricultural knowledge and modern AI technology. ' 
        'By providing actionable insights and diverse treatment options, it empowers farmers to make better ' 
        'decisions and improve crop yields sustainably.' 
    ) 
 
    # Save the document 
    doc.save('Project_Report.docx') 
    print("Project Report generated successfully as Project_Report.docx") 
 
if __name__ == "__main__": 
    try: 
        create_report() 
    except Exception as e: 
        print(f"Error: {e}") 
